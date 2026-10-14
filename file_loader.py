import base64
import csv
import os
from io import BytesIO

MAX_CHARS = 60_000 # around 15k tokens per file
MAX_IMAGE_EDGE = 1024 # longest side in pixels before downscaling
JPEG_QUALITY = 85

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".toml",
    ".json", ".yaml", ".yml", ".xml", ".html", ".htm", ".css",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".h", ".cpp", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".lua", ".r",
    ".sh", ".bat", ".ps1", ".sql", ".env",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}

MEDIA_TYPES = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/bmp",
}

class UnsupportedFile(Exception):
    pass

def trucate(text, name):
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + f"\n\n[...{name} truncated at {MAX_CHARS} chars...]"
    return text

def read_text(path):
    with open(path, "r", encoding="utf8", errors="replace") as f:
        return f.read()

def read_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise UnsupportedFile("pypdf not installed (pip install pypdf)")
    reader = PdfReader(path)
    return "\n\n".join(
        f"--- Page {i} ---\n{page.extract_text() or ''}"
        for i, page in enumerate(reader.pages, 1)
    )

def read_docx(path):
    try:
        import docx
    except ImportError:
        raise UnsupportedFile("python-docx not installed (pip install python-docx)")
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)

def read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        raise UnsupportedFile("openpyxl not installed (pip install openpyxl)")
    workbook = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            parts.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(parts)

def read_csv(path):
    with open(path, "r", encoding="utf8", errors="replace", newline="") as f:
        return "\n".join(" | ".join(row) for row in csv.reader(f))

def encode_image(image):
    from PIL import Image
    if max(image.size) > MAX_IMAGE_EDGE:
        scale = MAX_IMAGE_EDGE / max(image.size)
        image = image.resize(
            (max(1, round(image.width * scale)),
             max(1, round(image.height * scale))),
            Image.LANCZOS
        )

    has_alpha = image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info)

    buffer = BytesIO()
    if has_alpha:
        image.convert("RGBA").save(buffer, format="PNG", optimize=True)
        media_type = "image/png"
    else:
        image.convert("RGB").save(buffer, format="JPEG", quality=JPEG_QUALITY, optimize=True)
        media_type = "image/jpeg"

    return base64.b64encode(buffer.getvalue()).decode("ascii"), media_type

def read_image(path, ext):
    try:
        from PIL import Image
    except ImportError:
        with open(path, "rb") as f:
            return base64.b64encode(f.read()).decode("ascii"), MEDIA_TYPES[ext]

    with Image.open(path) as source:
        source.load() # animated GIF: this grabs the first frame
        return encode_image(source)

def load_image_bytes(raw, name="pasted.png"):
    from PIL import Image
    with Image.open(BytesIO(raw)) as source:
        source.load()
        data, media_type = encode_image(source)

    return {"name": name, "kind": "image", "data": data, "media_type": media_type}

def load_file(path):
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()

    if ext in IMAGE_EXTENSIONS:
       data, media_type = read_image(path, ext)
       return {"name": name, "kind": "image", "data": data,
                "media_type": media_type}

    if ext == ".pdf":
        text = read_pdf(path)
    elif ext == ".docx":
        text = read_docx(path)
    elif ext in (".xlsx", ".xlsm"):
        text = read_xlsx(path)
    elif ext == ".csv":
        text = read_csv(path)
    elif ext in TEXT_EXTENSIONS or ext == "":
        text = read_text(path)
    else:
        raise UnsupportedFile(f"Unsupported file type: {ext or 'no extension'}")

    if not text.strip():
        raise UnsupportedFile(f"No readable text in {name} (scanned PDF? needs OCR)")

    return {"name": name, "kind": "text", "text": trucate(text, name)}